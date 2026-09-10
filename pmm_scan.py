"""Extract file references from an MMD .pmm project and resolve broken links."""
import json
import os
import re
import time

REF_EXTS = [b".pmx", b".pmd", b".x", b".vmd", b".vpd", b".wav", b".avi", b".fx", b".fxsub"]
_DRIVE_RE = re.compile(r"[A-Za-z]:\\")
KIND_BY_EXT = {
    ".pmx": "モデル(.pmx)",
    ".pmd": "モデル(.pmd)",
    ".x": "アクセサリ(.x)",
    ".vmd": "モーション(.vmd)",
    ".vpd": "ポーズ(.vpd)",
    ".wav": "音声(.wav)",
    ".avi": "背景動画(.avi)",
    ".fx": "エフェクト(.fx)",
    ".fxsub": "エフェクト(.fxsub)",
}
INDEX_CACHE_NAME = "file_index_cache.json"


def extract_pmm_refs(pmm_path):
    """Walk each extension match backward to the nearest null/control byte,
    decode as cp932, then trim to the last drive-letter prefix (X:\\) found
    inside the run so any leading junk bytes (previous field's tail) are
    discarded. Longer matches make shorter ones redundant; only keep
    strings that were not fully contained in a longer one.

    Returns them in PMM byte order (the order objects were actually
    written to the file), not alphabetically — callers that fall back to
    positional matching against MMD's own load-order dialogs (see
    mmd_dialogs.run_autoload's model_fallback_queue) depend on this.
    """
    with open(pmm_path, "rb") as f:
        data = f.read()

    raw_hits = {}  # decoded string -> earliest byte offset seen at
    for ext in REF_EXTS:
        start_search = 0
        while True:
            idx = data.find(ext, start_search)
            if idx == -1:
                break
            end = idx + len(ext)
            i = idx
            while i > 0:
                b = data[i - 1]
                if b == 0x00 or (b < 0x20 and b not in (0x09,)):
                    break
                i -= 1
            raw = data[i:end]
            try:
                decoded = raw.decode("cp932")
            except Exception:
                decoded = None
            if decoded:
                matches = list(_DRIVE_RE.finditer(decoded))
                if matches:
                    decoded = decoded[matches[-1].start():]
                    if decoded not in raw_hits or i < raw_hits[decoded]:
                        raw_hits[decoded] = i
            start_search = end

    found = {}
    for h, offset in raw_hits.items():
        if any(h != other and h in other for other in raw_hits):
            continue
        found[h] = offset
    return [h for h, _off in sorted(found.items(), key=lambda kv: kv[1])]


def build_file_index(roots, cache_dir, force_rebuild=False, on_progress=None):
    """on_progress(root, dirs_seen, files_seen), called periodically (not per
    file) while walking, so a slow/large drive doesn't look like it's frozen.
    """
    cache_path = os.path.join(cache_dir, INDEX_CACHE_NAME)
    if not force_rebuild and os.path.exists(cache_path):
        with open(cache_path, "r", encoding="utf-8") as f:
            cache = json.load(f)
        if sorted(cache["roots"]) == sorted(roots):
            return cache["index"], cache["built_at"], cache["roots"]

    index = {}
    t0 = time.time()
    dirs_seen = 0
    files_seen = 0
    last_report = t0
    for root in roots:
        if not os.path.isdir(root):
            continue
        for dirpath, _dirnames, filenames in os.walk(root):
            dirs_seen += 1
            for name in filenames:
                files_seen += 1
                key = name.lower()
                full = os.path.join(dirpath, name)
                index.setdefault(key, []).append(full)
            if on_progress and time.time() - last_report >= 1.0:
                on_progress(root, dirs_seen, files_seen)
                last_report = time.time()
    if on_progress:
        on_progress(None, dirs_seen, files_seen)
    built_at = time.strftime("%Y-%m-%d %H:%M:%S")
    with open(cache_path, "w", encoding="utf-8") as f:
        json.dump({"index": index, "built_at": built_at, "roots": roots,
                   "seconds": round(time.time() - t0, 1)}, f, ensure_ascii=False)
    return index, built_at, roots


def _path_similarity(old_path, candidate_path):
    old_parts = [p.lower() for p in old_path.replace("/", "\\").split("\\")[:-1]]
    cand_parts = [p.lower() for p in candidate_path.replace("/", "\\").split("\\")[:-1]]
    score = 0
    for a, b in zip(reversed(old_parts), reversed(cand_parts)):
        if a == b:
            score += 1
        else:
            break
    return score


def _pick_by_date(candidates, pmm_mtime):
    """Tie-break same-similarity-score candidates by modification time: the
    PMM can only ever have referenced a file that already existed when it
    was saved, so prefer whichever tied candidate is the newest among those
    no newer than the PMM itself (closest match to what was actually there
    at the time). If every tied candidate is newer than the PMM (unusual —
    e.g. the whole folder got touched by a later copy/reorganize), fall
    back to the oldest one instead of guessing among files that all
    postdate the reference.
    """
    if pmm_mtime is None:
        return None
    dated = []
    for c in candidates:
        try:
            dated.append((os.path.getmtime(c), c))
        except OSError:
            continue
    if not dated:
        return None
    older_or_equal = [d for d in dated if d[0] <= pmm_mtime]
    if older_or_equal:
        return max(older_or_equal, key=lambda d: d[0])[1]
    return min(dated, key=lambda d: d[0])[1]


def _resolve_path(ref, index, pmm_mtime=None):
    """Core resolution logic shared by resolve_refs (pmm-embedded refs) and
    resolve_emm_refs (MMEffect assignment paths): does this path exist as
    given, and if not, what does the file index suggest instead? Returns
    a dict with exists/resolved_path/status/candidates — everything an
    entry needs except caller-specific identity fields (name/kind for a
    pmm ref, section/key for an emm entry)."""
    basename = os.path.basename(ref)
    exists = os.path.exists(ref)
    entry = {
        "exists": exists,
        "resolved_path": ref if exists else None,
        "status": "OK（元のパスに存在）" if exists else None,
        "candidates": [],
    }
    if not exists:
        candidates = index.get(basename.lower(), [])
        candidates = [c for c in candidates if c.lower() != ref.lower()]
        entry["candidates"] = candidates
        if len(candidates) == 1:
            entry["resolved_path"] = candidates[0]
            entry["status"] = "対応済み（自動検出・候補1件）"
        elif len(candidates) > 1:
            best_score = max(_path_similarity(ref, c) for c in candidates)
            tied = [c for c in candidates if _path_similarity(ref, c) == best_score]
            if best_score > 0 and len(tied) == 1:
                entry["resolved_path"] = tied[0]
                entry["status"] = f"対応済み（自動検出・候補{len(candidates)}件から最有力を選択）"
            elif best_score > 0 and len(tied) > 1:
                picked = _pick_by_date(tied, pmm_mtime)
                if picked:
                    entry["resolved_path"] = picked
                    entry["status"] = (
                        f"対応済み（候補{len(candidates)}件中{len(tied)}件が同点、"
                        "PMMの日付に近いものを選択）"
                    )
                else:
                    entry["status"] = f"要確認（候補{len(candidates)}件、絞り込めず）"
            else:
                entry["status"] = f"要確認（候補{len(candidates)}件、絞り込めず）"
        else:
            entry["status"] = "未対応（候補見つからず）"
    return entry


def resolve_refs(refs, index, pmm_mtime=None):
    results = []
    for ref in refs:
        basename = os.path.basename(ref)
        ext = os.path.splitext(basename)[1].lower()
        kind = KIND_BY_EXT.get(ext, ext)
        resolved = _resolve_path(ref, index, pmm_mtime)
        entry = {
            "name": os.path.splitext(basename)[0],
            "kind": kind,
            "old_path": ref,
        }
        entry.update(resolved)
        results.append(entry)
    return results


def parse_emm(emm_path):
    """Parse a .emm MMEffect assignment file (plain 'Key = Value' lines
    under '[Section]' headers, as written by MMEffect's own File > 設定を
    保存 command) into a flat list of (section, key, value) tuples. Same
    cp932 encoding as pmm-embedded strings. Values that aren't paths
    ('none', 'true'/'false', an object id like 'Acs1' for Owner=) are
    returned as-is — filtering to path-like values is the caller's job.
    """
    with open(emm_path, "rb") as f:
        data = f.read()
    text = data.decode("cp932", errors="replace")
    section = None
    entries = []
    for line in text.splitlines():
        line = line.strip()
        if not line or line.startswith(";"):
            continue
        if line.startswith("[") and line.endswith("]"):
            section = line[1:-1]
            continue
        if section is None or "=" not in line:
            continue
        key, _, value = line.partition("=")
        entries.append((section, key.strip(), value.strip()))
    return entries


def resolve_emm_refs(emm_path, index, pmm_mtime=None):
    """Check every .fx/.fxsub path an MMEffect assignment export actually
    references, and try to resolve any that are missing — same approach
    as resolve_refs, via the shared _resolve_path. [Object] entries are
    skipped: those are the accessory/model files MME's assignment dialog
    is driving, which pmm_fixer's own pmm scan already covers; this is
    specifically for the effect files themselves, which (per investigation
    on real projects) are often not stored as plain paths in the pmm at
    all, so they need to come from MME's own exported state instead."""
    results = []
    for section, key, value in parse_emm(emm_path):
        if not value.lower().endswith((".fx", ".fxsub")):
            continue
        resolved = _resolve_path(value, index, pmm_mtime)
        entry = {"section": section, "key": key, "old_path": value}
        entry.update(resolved)
        results.append(entry)
    return results


def write_report_xlsx(pmm_path, results, index_built_at, roots, out_path):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "リンク切れ一覧"

    FONT_NAME = "Arial"
    header_font = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    title_font = Font(name=FONT_NAME, bold=True, size=14)
    sub_font = Font(name=FONT_NAME, italic=True, size=10, color="666666")
    normal_font = Font(name=FONT_NAME, size=10)
    ok_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    ok_font = Font(name=FONT_NAME, size=10, color="006100")
    warn_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    warn_font = Font(name=FONT_NAME, size=10, color="9C6500")
    ng_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    ng_font = Font(name=FONT_NAME, size=10, color="9C0006")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells("A1:F1")
    ws["A1"] = f"{os.path.basename(pmm_path)}  リンク切れ一覧"
    ws["A1"].font = title_font

    ws.merge_cells("A2:F2")
    ws["A2"] = f"検索対象ドライブ: {', '.join(roots)}　/　インデックス作成: {index_built_at}"
    ws["A2"].font = sub_font

    headers = ["名称", "種類", "旧パス（PMM記載）", "解決パス／候補", "状態", "備考"]
    header_row = 4
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border

    r_idx = header_row + 1
    for entry in results:
        if entry["resolved_path"]:
            resolved_display = entry["resolved_path"]
        elif entry["candidates"]:
            resolved_display = "\n".join(entry["candidates"][:5])
        else:
            resolved_display = "(候補なし)"

        note = ""
        if entry["exists"]:
            note = "元のパスのまま存在"
        elif entry["resolved_path"] and entry["candidates"]:
            note = f"候補{len(entry['candidates'])}件中から自動選択"
        elif not entry["candidates"] and not entry["exists"]:
            note = "検索対象ドライブ内に同名ファイルなし"

        row = [entry["name"], entry["kind"], entry["old_path"], resolved_display, entry["status"], note]
        for c_idx, val in enumerate(row, start=1):
            c = ws.cell(row=r_idx, column=c_idx, value=val)
            c.font = normal_font
            c.border = border
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if c_idx == 5:
                if entry["exists"] or (entry["resolved_path"] and "自動検出" in entry["status"]):
                    c.font, c.fill = ok_font, ok_fill
                elif entry["resolved_path"]:
                    c.font, c.fill = ok_font, ok_fill
                elif entry["candidates"]:
                    c.font, c.fill = warn_font, warn_fill
                else:
                    c.font, c.fill = ng_font, ng_fill
                c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[r_idx].height = 34
        r_idx += 1

    widths = {"A": 26, "B": 16, "C": 46, "D": 46, "E": 30, "F": 30}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.row_dimensions[header_row].height = 24
    ws.freeze_panes = "A5"

    wb.save(out_path)
    return out_path


def write_effect_report_xlsx(pmm_path, results, out_path):
    """Same look as write_report_xlsx, but for resolve_emm_refs results
    (section/key/old_path instead of name/kind/old_path — MMEffect
    assignments don't have a single meaningful 'name' the way a pmm ref
    does, so the sheet is keyed by which assignment slot it came from)."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MMEffect状態"

    FONT_NAME = "Arial"
    header_font = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    title_font = Font(name=FONT_NAME, bold=True, size=14)
    sub_font = Font(name=FONT_NAME, italic=True, size=10, color="666666")
    normal_font = Font(name=FONT_NAME, size=10)
    ok_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    ok_font = Font(name=FONT_NAME, size=10, color="006100")
    warn_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    warn_font = Font(name=FONT_NAME, size=10, color="9C6500")
    ng_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    ng_font = Font(name=FONT_NAME, size=10, color="9C0006")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells("A1:E1")
    ws["A1"] = f"{os.path.basename(pmm_path)}  MMEffect割り当て状態"
    ws["A1"].font = title_font

    ws.merge_cells("A2:E2")
    ws["A2"] = (
        "MMEffectの「エフェクト割り当て」ダイアログから読み取った、実際に使われている.fx/.fxsubファイルの一覧です。"
        "自動での差し替えは行っていません（候補が複数ある場合に誤ったものを割り当てるリスクがあるため）。"
    )
    ws["A2"].font = sub_font
    ws["A2"].alignment = Alignment(wrap_text=True)

    headers = ["セクション", "キー", "エフェクトファイルのパス", "解決パス／候補", "状態"]
    header_row = 4
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border

    r_idx = header_row + 1
    for entry in results:
        if entry["resolved_path"]:
            resolved_display = entry["resolved_path"]
        elif entry["candidates"]:
            resolved_display = "\n".join(entry["candidates"][:5])
        else:
            resolved_display = "(候補なし)"

        row = [entry["section"], entry["key"], entry["old_path"], resolved_display, entry["status"]]
        for c_idx, val in enumerate(row, start=1):
            c = ws.cell(row=r_idx, column=c_idx, value=val)
            c.font = normal_font
            c.border = border
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if c_idx == 5:
                if entry["exists"] or entry["resolved_path"]:
                    c.font, c.fill = ok_font, ok_fill
                elif entry["candidates"]:
                    c.font, c.fill = warn_font, warn_fill
                else:
                    c.font, c.fill = ng_font, ng_fill
                c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[r_idx].height = 34
        r_idx += 1

    widths = {"A": 24, "B": 16, "C": 46, "D": 46, "E": 30}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.row_dimensions[header_row].height = 24
    ws.freeze_panes = "A5"

    wb.save(out_path)
    return out_path


def _write_plain_listing(title, headers, rows, out_path):
    """Tab-separated plain text, openable in Notepad on any Windows PC with
    zero dependencies — a fallback for anyone without Excel (or another
    .xlsx-capable app like LibreOffice) installed. UTF-8 with a BOM so
    older Notepad builds detect the encoding correctly instead of
    misreading it as the system ANSI codepage."""
    lines = [title, "=" * len(title), "", "\t".join(headers)]
    for row in rows:
        lines.append("\t".join("" if v is None else str(v) for v in row))
    with open(out_path, "w", encoding="utf-8-sig") as f:
        f.write("\n".join(lines) + "\n")
    return out_path


def _resolved_display(entry):
    if entry["resolved_path"]:
        return entry["resolved_path"]
    if entry["candidates"]:
        return " / ".join(entry["candidates"][:5])
    return "(候補なし)"


def write_plain_listing_for_refs(pmm_path, results, out_path):
    """Plain-text companion to write_report_xlsx (same resolve_refs results)."""
    headers = ["名称", "種類", "旧パス（PMM記載）", "解決パス／候補", "状態"]
    rows = [
        [entry["name"], entry["kind"], entry["old_path"], _resolved_display(entry), entry["status"]]
        for entry in results
    ]
    return _write_plain_listing(f"{os.path.basename(pmm_path)}  リンク切れ一覧", headers, rows, out_path)


def write_plain_listing_for_emm(pmm_path, results, out_path):
    """Plain-text companion to write_effect_report_xlsx (same resolve_emm_refs results)."""
    headers = ["セクション", "キー", "エフェクトファイルのパス", "解決パス／候補", "状態"]
    rows = [
        [entry["section"], entry["key"], entry["old_path"], _resolved_display(entry), entry["status"]]
        for entry in results
    ]
    return _write_plain_listing(f"{os.path.basename(pmm_path)}  MMEffect割り当て状態", headers, rows, out_path)
